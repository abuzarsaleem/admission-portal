from openpyxl import load_workbook

path = r"C:\Users\Ali Shan\Downloads\ADM-F000_Extracted_Tables_and_Sample_Data_v2.xlsx"
wb = load_workbook(path, data_only=True)
for name in [
    "admissions_general_criteria",
    "admissions_intakes",
    "admissions_programme_offerings",
]:
    ws = wb[name]
    print("\n" + "=" * 80)
    print(f"SHEET: {name}")
    print("=" * 80)
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        vals = ["" if v is None else str(v) for v in row]
        if all(v.strip() == "" for v in vals):
            continue
        print(f"{i:03d}| " + " | ".join(vals))
